# 빈 양식(엑셀) 생성 — 총입출금내역관리와 같은 열 + 헌금자명단 + 계정과목 참고
# renderer/defaults.js 의 계정과목·헌금종류를 그대로 사용
import json, re, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

root, out = sys.argv[1], sys.argv[2]
src = open(f"{root}/renderer/defaults.js", encoding="utf-8").read()
D = json.loads(re.search(r"=\s*(\{.*\});", src, re.S).group(1))
ACC = D["accounts"]
YEAR = datetime.date.today().year
F = "맑은 고딕"
base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
head_fill, head_font = PatternFill("solid", fgColor="27466E"), Font(name=F, size=10, bold=True, color="FFFFFF")
input_fill = PatternFill("solid", fgColor="FFFBE6")
ex_font = Font(name=F, size=10, italic=True, color="808080")
thin = Side(style="thin", color="C9D3CC"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
WON = '#,##0;-#,##0;"-"'
wb = Workbook()

def sheet(title, headers, widths):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(1, i, h); c.font = head_font; c.fill = head_fill; c.alignment = Alignment(horizontal="center"); c.border = box
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "A2"
    return ws

g = wb.active; g.title = "안내"
g.column_dimensions["A"].width = 3; g.column_dimensions["B"].width = 16; g.column_dimensions["C"].width = 100
g["B2"] = "성현교회 재정 앱 — 빈 양식"; g["B2"].font = Font(name=F, size=15, bold=True, color="27466E")
lines = [
    ("먼저", "이미 쓰고 계신 '총입출금내역관리.xlsx'와 '2026 수입예산'·'2026 지출예산' 파일은 그대로 앱에 올리면 됩니다. 이 양식은 새 해나 새 파일을 시작할 때 쓰세요."),
    ("입출금", "총입출금내역관리와 같은 열입니다. 거래 한 건 = 한 줄. 노란 칸에 입력, 회색 기울임 줄은 예시(가져올 때 제외)."),
    ("", "• 오프라인 헌금: 내용=이름, 입금=금액, 비고=오프라인, 헌금종류, Month=그 주일 날짜"),
    ("", "• 온라인 입금: 거래일시(통장 그대로), 내용, 입금, 비고=온라인, 헌금종류, Month"),
    ("", "• 지출: 거래일시, 내용, 출금, 지출계정과목(지출예산 항목 이름 그대로), Month"),
    ("", "• 현금입금: 오프라인 헌금을 통장에 넣은 금액은 '입금'이 아니라 '현금입금' 칸에 (이중 계산 방지)"),
    ("헌금종류", "목록에서 고릅니다. 어떤 목으로 들어가는지는 '계정과목' 시트 참고. 앞뒤 공백은 앱이 자동으로 지웁니다."),
    ("헌금자명단", "선택. 번호 순서가 주간 집계표 순서가 되고, 주소·주민등록번호는 기부금영수증에만 쓰입니다(이 PC에만 저장)."),
]
for r, (k, v) in enumerate(lines, 4):
    g.cell(r, 2, k).font = bold; c = g.cell(r, 3, v); c.font = base; c.alignment = Alignment(wrap_text=True)
g.sheet_view.showGridLines = False

H = ["거래일시", "내용", "출금", "입금", "거래후 잔액", "지출계정과목(적요)", "비고", "헌금종류", "Month", "현금입금"]
ws = sheet("입출금", H, [20, 22, 12, 12, 14, 20, 10, 14, 12, 12])
ex = [
    [None, "홍길동", None, 30000, None, None, "오프라인", "감사", datetime.date(YEAR, 1, 4), None],
    [f"{YEAR}.01.03 20:14:12", "홍길동십일조", None, 100000, 1000000, None, "온라인", "십일조", datetime.date(YEAR, 1, 4), None],
    [f"{YEAR}.01.05 12:36:02", "1/4 주일헌금", None, None, 1030000, None, "현금입금", None, datetime.date(YEAR, 1, 4), 30000],
    [f"{YEAR}.01.06 09:10:00", "1월 전기료", 450000, None, 580000, "전기요금", None, None, datetime.date(YEAR, 1, 11), None],
]
for r, row in enumerate(ex, 2):
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v); cell.font = ex_font; cell.border = box
        if c in (3, 4, 5, 10): cell.number_format = WON
        if c == 9: cell.number_format = "yyyy-mm-dd"
    ws.cell(r, 11, "예시").font = ex_font
for r in range(6, 3006):
    for c in range(1, 11):
        cell = ws.cell(r, c); cell.font = base; cell.border = box; cell.fill = input_fill
        if c in (3, 4, 5, 10): cell.number_format = WON
        if c == 9: cell.number_format = "yyyy-mm-dd"

kinds = [k for a in ACC for k in a["kinds"]]
exp = [a["hang"] for a in ACC if a["type"] == "지출"]
wa = sheet("계정과목", ["구분", "관", "항", "번호", "목 / 항목", "연결된 헌금종류", "비고"], [8, 14, 14, 7, 24, 50, 26])
for i, a in enumerate(ACC, 2):
    for j, v in enumerate([a["type"], a["gwan"], a["mid"], a["code"], a["hang"], ", ".join(a["kinds"]), a["note"]], 1):
        c = wa.cell(i, j, v); c.font = base; c.border = box
wa["I1"] = "헌금종류"; wa["J1"] = "지출계정과목"
for i, v in enumerate(kinds, 2): wa.cell(i, 9, v)
for i, v in enumerate(exp, 2): wa.cell(i, 10, v)
wa.column_dimensions["I"].hidden = True; wa.column_dimensions["J"].hidden = True
for rng, f in [("H6:H3005", f"=계정과목!$I$2:$I${len(kinds)+1}"), ("F6:F3005", f"=계정과목!$J$2:$J${len(exp)+1}"), ("G6:G3005", '"오프라인,온라인,현금입금"')]:
    dv = DataValidation(type="list", formula1=f, allow_blank=True, showErrorMessage=False); ws.add_data_validation(dv); dv.add(rng)

wd = sheet("헌금자명단", ["번호", "이름", "직분", "연락처", "주소(영수증용)", "주민등록번호(영수증용·선택)", "비고"], [7, 12, 10, 15, 40, 24, 22])
for j, v in enumerate([1, "홍길동", "집사", "010-1234-5678", "서울특별시 동작구 ○○로 00", "", "예시"], 1):
    c = wd.cell(2, j, v); c.font = ex_font; c.border = box
for r in range(3, 150):
    for c in range(1, 8): cell = wd.cell(r, c); cell.font = base; cell.border = box; cell.fill = input_fill
wb.move_sheet("계정과목", offset=1)
wb.save(out); print("ok", out)
